import argparse
import re
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter


DISEASE_ORDER = ["AMD", "Glaucoma", "Health", "ALL"]

MODEL_ORDER = [
    "dinov3-vits16",
    "dinov3-vitl16",
    "dinov3-vitb16",
    "RETFound-dinov2-meh",
    "RETFound-dinov2-shanghai",
    "FLAIR",
    "visionFM",
    "EyeCLIP",
    "RetiZero",
]


def canonical_model(raw_model: str):
    s = str(raw_model).strip()
    lower = s.lower()

    mapping = {
        "retfound": "RETFound",
        "dinov3_vits16": "dinov3-vits16",
        "dinov3_vitl16": "dinov3-vitl16",
        "dinov3_vitb16": "dinov3-vitb16",
        "retfound_dinov2_meh": "RETFound-dinov2-meh",
        "retfound-dinov2-meh": "RETFound-dinov2-meh",
        "retfound_dinov2_shanghai": "RETFound-dinov2-shanghai",
        "retfound-dinov2-shanghai": "RETFound-dinov2-shanghai",
        "flair": "FLAIR",
        "visionfm": "visionFM",
        "eyeclip": "EyeCLIP",
        "retizero": "RetiZero",
    }

    key = lower.replace("-", "_")
    return mapping.get(key, mapping.get(lower, s))


def fmt_value(mean, std):
    if pd.isna(mean) or pd.isna(std):
        return "—"
    return f"{mean:.4f} ± {std:.4f}"


def safe_sheet_name(name):
    name = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(name))
    return name[:31]


parser = argparse.ArgumentParser()
parser.add_argument("--input_csv", required=True)
parser.add_argument("--output_xlsx", required=True)
args = parser.parse_args()

df = pd.read_csv(args.input_csv)

required = {"model", "metric", "disease", "mean", "std"}
missing = required - set(df.columns)
if missing:
    raise ValueError(f"CSV 缺少必要列: {missing}")

df["model_clean"] = df["model"].apply(canonical_model)

extra_diseases = [d for d in df["disease"].dropna().unique() if d not in DISEASE_ORDER]
disease_order = DISEASE_ORDER + extra_diseases

wb = Workbook()
wb.remove(wb.active)

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for metric in df["metric"].dropna().unique():
    sub = df[df["metric"] == metric].copy()
    ws = wb.create_sheet(safe_sheet_name(metric))

    ws.cell(1, 1, f"external validation - {metric}")
    ws.cell(1, 1).font = Font(bold=True, size=12)

    headers = ["model"] + disease_order
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(2, col, h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border

    models_in_csv = list(dict.fromkeys(sub["model_clean"]))
    model_order = MODEL_ORDER + [m for m in models_in_csv if m not in MODEL_ORDER]

    row = 3

    for model in model_order:
        sub_m = sub[sub["model_clean"] == model]
        if sub_m.empty:
            continue

        ws.cell(row, 1, model)

        for col, disease in enumerate(disease_order, start=2):
            hit = sub_m[sub_m["disease"] == disease]

            if hit.empty:
                value = "—"
            else:
                r = hit.iloc[0]
                value = fmt_value(r["mean"], r["std"])

            ws.cell(row, col, value)

        row += 1

    # 全部左对齐 + 边框
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in r:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if cell.row >= 2:
                cell.border = border

    # 列宽
    ws.column_dimensions["A"].width = 28
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws.freeze_panes = "B3"

wb.save(args.output_xlsx)
print(f"Saved to: {Path(args.output_xlsx).resolve()}")
